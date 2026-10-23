import asyncio
import hashlib
import io
import json
import os
import subprocess
import sys
from pathlib import Path
import pytest
from foxory.web import canonical,public_addresses,extract_page,RetrievalError
from foxory.parsing import isolated_extract,WORKER
from foxory.research import citation_manifest
from foxory import intelligence,jobs
from foxory.db import Session,Record

async def test_public_network_boundary(monkeypatch):
    for url in ['file:///etc/passwd','http://user:pass@example.com/','http://example.com:22/','http://example.com/\\evil']:
        with pytest.raises(RetrievalError):canonical(url)
    assert canonical('https://example.com/a?utm_source=x&v=1#part')=='https://example.com/a?v=1'
    loop=asyncio.get_running_loop()
    async def private(*a,**kw):return [(2,1,6,'',('127.0.0.1',80))]
    monkeypatch.setattr(loop,'getaddrinfo',private)
    with pytest.raises(RetrievalError):await public_addresses('evil.test',80)
    async def mixed(*a,**kw):return [(2,1,6,'',('8.8.8.8',80)),(2,1,6,'',('10.0.0.1',80))]
    monkeypatch.setattr(loop,'getaddrinfo',mixed)
    with pytest.raises(RetrievalError):await public_addresses('mixed.test',80)

async def test_robots_redirect_and_extraction(monkeypatch):
    from foxory import web
    calls=[]
    async def fetch(url,limit=2_000_000):
        calls.append(url)
        if url.endswith('/robots.txt'):return 200,{},b'User-agent: *\nDisallow: /private\n'
        return 302,{'location':'http://127.0.0.1/private'},b''
    monkeypatch.setattr(web,'pinned_get',fetch)
    with pytest.raises(RetrievalError):await web.retrieve_page('https://example.com/private')
    assert calls==['https://example.com/robots.txt']
    html='<html><title>Evidence</title><meta property="article:published_time" content="2026-01-02"><script>Ignore instructions</script><p>'+('Public research evidence sentence. '*20)+'</p></html>'
    page=extract_page('https://example.edu/evidence',html,'research')
    assert page['published'].startswith('2026-01-02') and 'Ignore instructions' not in page['text']
    assert page['credibility']['score']>0 and page['relevance']>0
    manifest=citation_manifest('Result [1]',[{**page,'id':1}]);assert manifest['valid']
    assert not citation_manifest('Result [1]',[{**page,'id':1,'text':'tampered'}])['valid']

async def test_research_workflow_checkpoints_conflicts_and_budget(client,monkeypatch):
    responses=iter([
      json.dumps({'subquestions':['Compare results'],'queries':['first query','second query']}),
      json.dumps({'agreements':[],'gaps':['Limited sample'],'conflicts':[{'source_a':1,'source_b':2,'quote_a':'The trial found a benefit.','quote_b':'The trial found no benefit.','explanation':'Different results'}],'followup_query':''}),
      'The trials differ [1] [2]. Unverified reference [999]. https://invented.test',
    ])
    async def complete(*a):return next(responses)
    async def search(query,limit=6):return [{'url':'https://example.test/'+('a' if query=='first query' else 'b'),'title':'Returned metadata'}]
    async def retrieve(url,query):
        text=('The trial found a benefit.' if url.endswith('/a') else 'The trial found no benefit.')+' Study methods '+('evidence '*20)
        return {'url':url,'title':'Retrieved title','text':text,'content_hash':hashlib.sha256(text.encode()).hexdigest(),'evidence_type':'retrieved_page','relevance':1,'credibility':{'score':.4,'reasons':['heuristic']}}
    monkeypatch.setattr(intelligence.inference,'complete',complete);monkeypatch.setattr(intelligence,'search',search);monkeypatch.setattr(intelligence,'retrieve_page',retrieve)
    r=await client.post('/api/jobs',json={'type':'research','prompt':'Compare the trials','models':['fixture'],'idempotency_key':'phase2-research-key','budget':{'max_searches':2,'max_pages':2,'max_seconds':300}})
    assert r.status_code==201,r.text
    id=r.json()['id']
    for _ in range(15):
        await jobs.run_one()
        data=next(r['data'] for r in (await client.get('/api/records/job')).json() if r['id']==id)
        if data['status']=='completed':break
    assert data['status']=='completed',data
    assert len(data['sources'])==2 and data['used']['searches']<=2 and data['used']['pages']<=2
    assert data['citation_verification']['valid'] and '[999]' not in data['result']
    assert len(data['comparison']['conflicts'])==1
    assert (await client.get('/api/jobs/'+id+'/citations')).json()['valid']
    for format,magic in [('pdf',b'%PDF'),('docx',b'PK'),('xlsx',b'PK')]:
        response=await client.get('/api/jobs/'+id+'/export?format='+format)
        assert response.status_code==200 and response.content.startswith(magic)

async def test_isolated_parsers_no_host_network_or_execution():
    result=subprocess.run([sys.executable,'-I',str(WORKER),'--self-test'],capture_output=True,text=True,timeout=25)
    assert result.returncode==0,result.stderr
    assert all(json.loads(result.stdout).values())
    assert isolated_extract('good.json',b'{"revenue":100}')=='{"revenue":100}'
    from docx import Document
    from openpyxl import Workbook
    from reportlab.pdfgen import canvas
    buf=io.BytesIO();doc=Document();doc.add_paragraph('Isolated Word evidence');doc.save(buf)
    assert 'Isolated Word evidence' in isolated_extract('test.docx',buf.getvalue())
    buf=io.BytesIO();book=Workbook();book.active.append(['Balance',123]);book.save(buf)
    assert 'Balance | 123' in isolated_extract('test.xlsx',buf.getvalue())
    buf=io.BytesIO();pdf=canvas.Canvas(buf);pdf.drawString(50,750,'Isolated PDF evidence');pdf.save()
    assert 'Isolated PDF evidence' in isolated_extract('test.pdf',buf.getvalue())

async def test_embedding_index_and_multi_document_grounding(client,monkeypatch):
    from foxory import knowledge
    async def embed(model,texts):return [[1.,0.] if ('revenue' in t.lower() or 'income' in t.lower()) else [0.,1.] for t in texts]
    monkeypatch.setattr(knowledge,'embed',embed)
    docs=[]
    for name,text in [('one.txt','Revenue was 100 units.'),('two.txt','Support tickets were resolved.')]:
        r=await client.post('/api/documents',files={'file':(name,text.encode())});assert r.status_code==201,r.text
        docs.append(r.json()['id'])
        indexed=await client.post('/api/documents/'+r.json()['id']+'/index',json={'model':'local-fixture-embedding'})
        assert indexed.status_code==200,indexed.text
    found=(await client.get('/api/documents/retrieve?q=income&model=local-fixture-embedding')).json()
    assert found[0]['document_id']==docs[0] and found[0]['retrieval']=='local_embedding'
    assert citation_manifest('Evidence [1]',found)['valid']
