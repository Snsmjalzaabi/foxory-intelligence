import asyncio
import io
import json
import zipfile
from pathlib import Path
import pytest
from sqlalchemy import select
from foxory.db import Session, Record, Message
from foxory import jobs
from foxory.main import app, active_chats
from foxory.documents import extract
from foxory.research import validate_citations
from fastapi import HTTPException

async def record(c,kind,title='Example',data=None):
    r=await c.post('/api/records/'+kind,json={'title':title,'data':data or {}})
    assert r.status_code==201,r.text
    return r.json()

async def test_auth_origin_and_host(client):
    assert (await client.get('/api/health',headers={'Authorization':''})).status_code==200
    assert (await client.get('/api/overview',headers={'Authorization':''})).status_code==401
    assert (await client.post('/api/records/memory',json={'title':'a'},headers={'Origin':'https://evil.example'})).status_code==403
    assert (await client.get('/api/overview',headers={'Host':'evil.example'})).status_code==400
    r=await client.get('/api/overview')
    assert r.status_code==200
    assert "frame-ancestors 'none'" in r.headers['content-security-policy']

async def test_crud_and_export(client):
    r=await record(client,'memory',data={'body':'User approved fact'})
    assert len((await client.get('/api/records/memory?q=approved')).json())==1
    assert (await client.put('/api/records/'+r['id'],json={'title':'Renamed','data':{'body':'changed'}})).status_code==200
    export=(await client.get('/api/workspace/export')).json()
    assert export['records'][0]['title']=='Renamed'
    assert (await client.delete('/api/records/'+r['id'])).status_code==204
    assert (await client.delete('/api/records/'+r['id'])).status_code==404

async def test_chat_persistence_regeneration_errors(client,monkeypatch):
    from foxory import inference
    async def stream(*args,**kwargs):
        yield {'message':{'content':'Hello '}}
        yield {'message':{'content':'local world'},'done':True,'eval_count':3}
    monkeypatch.setattr(inference,'stream',stream)
    r=await record(client,'conversation','Hello')
    body={'content':'Hello','model':'fixture-model'}
    result=await client.post('/api/conversations/'+r['id']+'/chat',json=body)
    assert 'event: done' in result.text and 'local world' in result.text
    messages=(await client.get('/api/conversations/'+r['id']+'/messages')).json()
    assert [m['role'] for m in messages]==['user','assistant']
    assert messages[1]['content']=='Hello local world'
    assert messages[1]['metadata']['eval_count']==3
    assert not active_chats
    await client.post('/api/conversations/'+r['id']+'/chat',json={**body,'regenerate':True})
    assert len((await client.get('/api/conversations/'+r['id']+'/messages')).json())==2
    assert 'Hello local world' in (await client.get('/api/conversations/'+r['id']+'/export')).text
    async def broken(*args,**kwargs):
        raise RuntimeError('sensitive service detail')
        yield
    monkeypatch.setattr(inference,'stream',broken)
    result=await client.post('/api/conversations/'+r['id']+'/chat',json=body)
    assert 'event: error' in result.text and 'sensitive service detail' not in result.text
    assert not active_chats
    await client.delete('/api/records/'+r['id'])
    with Session() as s: assert not s.scalars(select(Message)).all()

async def test_chat_busy_guard(client):
    r=await record(client,'conversation')
    active_chats.add(r['id'])
    assert (await client.post('/api/conversations/'+r['id']+'/chat',json={'content':'Hi','model':'m'})).status_code==409

async def test_upload_validation_and_retrieval(client):
    r=await client.post('/api/documents',files={'file':('../../memo.txt',b'Foxory revenue was 100 units in June.','text/plain')})
    assert r.status_code==201,r.text
    doc=r.json();assert doc['title']=='memo.txt'
    found=(await client.get('/api/documents/search?q=revenue')).json()
    assert found[0]['document_id']==doc['id']
    assert (await client.get('/api/documents/'+doc['id']+'/download')).content==b'Foxory revenue was 100 units in June.'
    for name,content in [('bad.exe',b'code'),('bad.json',b'{'),('binary.txt',b'\x00'),('bad.pdf',b'not a PDF')]:
        assert (await client.post('/api/documents',files={'file':(name,content)})).status_code in {415,422}
    assert (await client.post('/api/documents',files={'file':('big.txt',b'a'*(10*1024*1024+1))})).status_code==413
    await client.delete('/api/records/'+doc['id'])
    assert not (await client.get('/api/documents/search?q=revenue')).json()

async def test_archive_and_office_parsers(client):
    from docx import Document
    from openpyxl import Workbook
    buf=io.BytesIO();doc=Document();doc.add_paragraph('Quarterly planning');doc.save(buf)
    assert 'Quarterly planning' in extract('doc.docx',buf.getvalue())
    buf=io.BytesIO();book=Workbook();book.active.append(['Revenue',100]);book.save(buf)
    assert 'Revenue | 100' in extract('table.xlsx',buf.getvalue())
    buf=io.BytesIO()
    with zipfile.ZipFile(buf,'w') as z:z.writestr('word/vbaProject.bin',b'fake macro')
    with pytest.raises(HTTPException):extract('macro.docx',buf.getvalue())

async def test_lead_duplicate_and_csv_injection(client):
    r=await record(client,'lead','=SUM(1,2)',{'email':'owner@example.test','status':'new'})
    duplicate=await client.post('/api/records/lead',json={'title':'other','data':{'email':'OWNER@example.test'}})
    assert duplicate.status_code==409
    csv=(await client.get('/api/leads/export')).text
    assert "'=SUM(1,2)" in csv
    other=await record(client,'lead','Second',{'email':'second@example.test'})
    assert (await client.put('/api/records/'+other['id'],json={'title':'Changed','data':{'email':'owner@example.test'}})).status_code==409

async def test_approval_snapshot_enforcement(client):
    draft=await record(client,'draft','Reply',{'body':'Original authorized draft'})
    request=await client.post('/api/drafts/'+draft['id']+'/request-approval');assert request.status_code==200
    approval=request.json()
    again=(await client.post('/api/drafts/'+draft['id']+'/request-approval')).json();assert again['id']==approval['id']
    await client.put('/api/records/'+draft['id'],json={'title':'Edited','data':{'body':'Changed'}})
    r=await client.post('/api/approvals/'+approval['id']+'/approve')
    assert r.json()['data']['snapshot']['data']['body']=='Original authorized draft'
    assert r.json()['data']['external_execution'] is False
    assert (await client.post('/api/approvals/'+approval['id']+'/reject')).status_code==409
    assert (await client.put('/api/records/'+approval['id'],json={'title':'Bypass','data':{'status':'approved'}})).status_code==403
    assert (await client.post('/api/records/approval',json={'title':'Bypass'})).status_code==422
    assert (await client.delete('/api/records/'+approval['id'])).status_code==403

async def create_job(client,type='research',key='unique-key-123',models=None,**extra):
    r=await client.post('/api/jobs',json={'type':type,'prompt':'Research local AI','models':models or ['a'],'idempotency_key':key,**extra})
    assert r.status_code==201,r.text
    return r.json()

async def test_job_idempotency_and_transitions(client):
    first=await create_job(client)
    second=await create_job(client)
    assert first['id']==second['id']
    assert len((await client.get('/api/records/job')).json())==1
    assert (await client.post('/api/jobs/'+first['id']+'/resume')).status_code==409
    assert (await client.post('/api/jobs/'+first['id']+'/pause')).json()['data']['status']=='paused'
    assert (await client.post('/api/jobs/'+first['id']+'/resume')).json()['data']['status']=='queued'
    assert (await client.post('/api/jobs/'+first['id']+'/cancel')).json()['data']['status']=='cancelled'
    assert not await jobs.run_one()

async def test_research_checkpoints_and_cancellation(client,monkeypatch):
    async def search(q,limit=6):return [{'id':1,'title':'Fixture evidence','url':'https://example.test/a','snippet':'A fact','evidence_type':'search_snippet'}]
    async def complete(*a):return 'Evidence [1]. Unsupported [99]. https://fabricated.test'
    monkeypatch.setattr(jobs,'search',search);monkeypatch.setattr(jobs,'complete',complete)
    job=await create_job(client)
    assert await jobs.run_one()
    saved=(await client.get('/api/records/job')).json()[0]
    assert saved['data']['step']==1 and saved['data']['status']=='queued'
    await jobs.run_one();await jobs.run_one()
    done=(await client.get('/api/records/job')).json()[0]
    assert done['data']['status']=='completed'
    assert len(done['data']['sources'])==1
    assert '[99]' not in done['data']['result'] and 'fabricated.test' not in done['data']['result']
    assert 'https://example.test/a' in (await client.get('/api/jobs/'+job['id']+'/export')).text
    # Cancel an in-flight operation; verify coroutine cancellation, not merely discarded output.
    entered=asyncio.Event();cancelled=asyncio.Event()
    async def slow_search(*a):
        entered.set()
        try:await asyncio.Event().wait()
        finally:cancelled.set()
    monkeypatch.setattr(jobs,'search',slow_search)
    job=await create_job(client,key='another-key-123')
    runner=asyncio.create_task(jobs.run_one());await asyncio.wait_for(entered.wait(),2)
    await client.post('/api/jobs/'+job['id']+'/cancel')
    await asyncio.wait_for(runner,2)
    assert cancelled.is_set()
    saved=[r for r in (await client.get('/api/records/job')).json() if r['id']==job['id']][0]
    assert saved['data']['status']=='cancelled' and saved['data']['step']==0

async def test_council_sequential_and_failure_retry(client,monkeypatch):
    calls=[]
    async def complete(model,messages):calls.append(model);return 'Independent analysis' if len(calls)<3 else 'Consensus does not verify facts.'
    monkeypatch.setattr(jobs,'complete',complete)
    j=await create_job(client,'council',models=['small-a','small-b'])
    for _ in range(3):await jobs.run_one()
    assert calls==['small-a','small-b','small-a']
    done=(await client.get('/api/records/job')).json()[0]
    assert done['data']['status']=='completed' and len(done['data']['analyses'])==2
    async def broken(*args):raise RuntimeError('provider disconnected')
    monkeypatch.setattr(jobs,'search',broken)
    j=await create_job(client,key='failure-key-123')
    for attempt in range(3):
        await jobs.run_one()
        state=[r for r in (await client.get('/api/records/job')).json() if r['id']==j['id']][0]
        assert state['data']['status']=='failed'
        retry=await client.post('/api/jobs/'+j['id']+'/retry')
        assert retry.status_code==(200 if attempt<2 else 409)

async def test_agent_permission_boundary(client):
    assert (await client.post('/api/records/agent',json={'title':'Unsafe','data':{'tools':['shell']}})).status_code==422
    agent=await record(client,'agent','Finance',{'enabled':False,'model':'small','tools':[]})
    r=await client.post('/api/jobs',json={'type':'agent','prompt':'Budget','models':['small'],'agent_id':agent['id'],'idempotency_key':'agent-key-123'})
    assert r.status_code==409

async def test_request_limit_and_rate_limit(client):
    r=await client.post('/api/records/memory',content=b' '*(11*1024*1024))
    assert r.status_code==413
    from foxory.main import requests_by_client
    import time
    requests_by_client['127.0.0.1'].extend([time.monotonic()]*180)
    assert (await client.get('/api/overview')).status_code==429


def test_citation_filter():
    assert '[77]' not in validate_citations('Claim [77] [1]',[{'id':1}])
    assert 'evil.test' not in validate_citations('[invented](https://evil.test)',[])

async def test_restart_recovery_preserves_checkpoint(client,monkeypatch):
    async def search(*args):return [{'id':1,'url':'https://example.test/evidence','title':'Evidence','snippet':'fixture'}]
    async def complete(*args):return 'Final synthesis [1]'
    monkeypatch.setattr(jobs,'search',search);monkeypatch.setattr(jobs,'complete',complete)
    j=await create_job(client,key='restart-recovery-key')
    await jobs.run_one()
    with Session() as s:
        r=s.get(Record,j['id']);d=json.loads(r.payload);assert d['step']==1
        d['status']='running';r.payload=json.dumps(d);s.commit()
    worker=asyncio.create_task(jobs.worker())
    try:
        async def finished():
            while True:
                with Session() as s:data=json.loads(s.get(Record,j['id']).payload)
                if data['status']=='completed':return data
                await asyncio.sleep(.05)
        result=await asyncio.wait_for(finished(),3)
        assert result['step']==3 and len(result['sources'])==1
    finally:
        worker.cancel()
        with __import__('contextlib').suppress(asyncio.CancelledError):await worker

async def test_concurrent_duplicate_submissions(client):
    body={'title':'Same lead','data':{'email':'same@example.test'}}
    responses=await asyncio.gather(*[client.post('/api/records/lead',json=body) for _ in range(6)])
    assert sum(r.status_code==201 for r in responses)==1
    assert all(r.status_code in {201,409} for r in responses)
    jobs_created=await asyncio.gather(*[create_job(client,key='concurrent-idempotency-key') for _ in range(6)])
    assert len({r['id'] for r in jobs_created})==1
