"""Trusted report writers; generated text is plain data, never executable markup."""
import io
from html import escape
from docx import Document
from openpyxl import Workbook
from reportlab.platypus import SimpleDocTemplate,Paragraph,Spacer
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.lib.pagesizes import A4

def report(title:str,text:str,sources:list,format:str):
    refs=[f"[{s['id']}] {s['title']} — {s.get('url','')} | retrieved {s.get('retrieved_at','n/a')} | hash {s.get('content_hash','')}" for s in sources]
    lines=(text+'\n\nSources\n'+'\n'.join(refs)).splitlines()
    out=io.BytesIO()
    if format=='pdf':
        style=getSampleStyleSheet();story=[Paragraph(escape(title),style['Title']),Spacer(1,12)]
        for line in lines:
            for i in range(0,max(1,len(line)),1500):story.append(Paragraph(escape(line[i:i+1500]) or ' ',style['BodyText']));story.append(Spacer(1,5))
        SimpleDocTemplate(out,pagesize=A4).build(story);mime='application/pdf'
    elif format=='docx':
        doc=Document();doc.add_heading(title,0)
        for line in lines:doc.add_paragraph(line)
        doc.save(out);mime='application/vnd.openxmlformats-officedocument.wordprocessingml.document'
    elif format=='xlsx':
        book=Workbook();sheet=book.active;sheet.title='Report'
        def safe(v):
            v=str(v)[:32000]
            return "'"+v if v.startswith(('=','+','-','@','\t','\r')) else v
        sheet.append([safe(title)])
        for line in text.splitlines():sheet.append([safe(line)])
        evidence=book.create_sheet('Sources');evidence.append(['ID','Title','URL','Retrieved','Hash','Evidence'])
        for source in sources:evidence.append([safe(source.get(k,'')) for k in ['id','title','url','retrieved_at','content_hash','text']])
        book.save(out);mime='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
    else:raise ValueError('Unknown report format')
    return out.getvalue(),mime
