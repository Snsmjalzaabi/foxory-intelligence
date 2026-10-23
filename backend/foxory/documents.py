import csv
import io
import json
import zipfile
from pathlib import Path
from fastapi import HTTPException

EXTENSIONS = {".txt", ".md", ".csv", ".json", ".pdf", ".docx", ".xlsx"}

def extract(name: str, raw: bytes):
    ext = Path(name).suffix.lower()
    if ext not in EXTENSIONS: raise HTTPException(415, "Unsupported document format")
    try:
        if ext in {".docx", ".xlsx"}:
            with zipfile.ZipFile(io.BytesIO(raw)) as z:
                infos = z.infolist()
                if len(infos) > 2000 or sum(x.file_size for x in infos) > 50_000_000:
                    raise ValueError("Archive expansion limit exceeded")
                if any("vbaproject" in x.filename.casefold() for x in infos): raise ValueError("Macros are not accepted")
        if ext == ".pdf":
            from pypdf import PdfReader
            reader = PdfReader(io.BytesIO(raw))
            if reader.is_encrypted or len(reader.pages) > 300: raise ValueError("Encrypted or oversized PDF")
            text = "\n".join(p.extract_text() or "" for p in reader.pages)
        elif ext == ".docx":
            from docx import Document
            doc = Document(io.BytesIO(raw))
            text = "\n".join([p.text for p in doc.paragraphs] + [" | ".join(c.text for c in row.cells) for t in doc.tables for row in t.rows])
        elif ext == ".xlsx":
            from openpyxl import load_workbook
            wb = load_workbook(io.BytesIO(raw), read_only=True, data_only=True)
            lines = []
            for sheet in wb:
                lines.append(f"Sheet: {sheet.title}")
                for i, row in enumerate(sheet.values):
                    if i >= 10000: raise ValueError("Spreadsheet row limit exceeded")
                    lines.append(" | ".join(str(v) if v is not None else "" for v in row))
            wb.close()
            text = "\n".join(lines)
        else:
            text = raw.decode("utf-8-sig")
            if "\x00" in text: raise ValueError("Binary content rejected")
            if ext == ".json": json.loads(text)
            if ext == ".csv": list(csv.reader(io.StringIO(text)))
        if len(text) > 2_000_000: raise ValueError("Extracted text limit exceeded")
        return text
    except HTTPException: raise
    except Exception as e: raise HTTPException(422, f"Document rejected: {type(e).__name__}") from e

def retrieve(records, query, limit=6):
    terms = set(query.lower().split())
    candidates = []
    for r in records:
        data = json.loads(r.payload)
        text = data.get("text", "")
        for offset in range(0, len(text), 1000):
            chunk = text[offset:offset+1400]
            score = sum(chunk.lower().count(t) for t in terms if len(t)>2)
            if score: candidates.append(dict(document_id=r.id, title=r.title, offset=offset, text=chunk, score=score))
    return sorted(candidates, key=lambda x:x["score"], reverse=True)[:limit]
